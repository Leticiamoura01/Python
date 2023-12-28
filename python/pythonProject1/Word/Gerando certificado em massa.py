from docx import Document
from docx.shared import Pt

from openpyxl import load_workbook
import os

#Pegando o caminho do arquivo
nome_arquivo_alunos = "Aluno.xlsx"
planilhaDadosAlunos = load_workbook(nome_arquivo_alunos)

#Selecionando a planilha/sheet/aba
sheet_selecionada = planilhaDadosAlunos["Nomes"]

for linha in range(2, len(sheet_selecionada["A"]) + 1):

        #Abre o arquivo do word
        arquivoWord = Document("C:\\Users\kauan\Downloads\\1 rpa\python\\pythonProject1\\Word\\Certificado2.docx")

        #Seleciona o estilo
        estilo = arquivoWord.styles["Normal"]

        #Pegando o nome do aluno quando passar na celula
        nomeAluno = sheet_selecionada['A%s' % linha].value

        #for = para
        for paragrafo in arquivoWord.paragraphs:

            #if = se
            if "@nome" in paragrafo.text:
                paragrafo.text = nomeAluno
                fonte = estilo.font
                fonte.name = "Calibri (Corpo)"
                fonte.size = Pt(24)

        caminhoCertificado = "coloca o caminho dapasta criada so pra certificado"

        #Salvando o certificado do aluno
        arquivoWord.save(caminhoCertificado)

print("Certificados gerados com sucesso!")

